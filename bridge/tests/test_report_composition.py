from __future__ import annotations

import base64
import pathlib
import zipfile

import pytest
from docx import Document

from noc_bridge.docx_render import render_document
from noc_bridge.failures import EvidenceIntegrityError, EvidenceRetrievalError
from noc_bridge.report_composition import compose_report
from noc_bridge.report_document import BilingualText
from noc_bridge.validation import OutputValidationError, validate_output


ROOT = pathlib.Path(__file__).resolve().parents[2]
SKILLS_DIR = ROOT / "skills"

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _snapshot() -> dict:
    return {
        "shift_starts_at": "2026-09-14T00:00:00+00:00",
        "shift_ends_at": "2026-09-14T08:00:00+00:00",
        "shift_timezone": "Asia/Manila",
        "shift_display_name": "MS",
        "report_skill_snapshot_id": "report-v2",
        "report_skill_execution_hash": "report-exec-v2",
        "incidents": [
            {
                "id": "inc-001", "display_id": "INC-001", "title": "Payment timeout",
                "status": "RECOVERED", "service": "payments", "environment": "prod",
                "triggered_at": "2026-09-14T01:00:00+00:00", "recovered_at": "2026-09-14T01:30:00+00:00",
                "trigger_value": "95%", "teams_url": "https://teams.example/inc-001",
                "grafana_url": "https://grafana.example/inc-001", "log_filename": "payments.log",
                "screenshots": [{"bucket": "noc-evidence", "object_key": "inc-001.png", "filename": "alert.png"}],
                "analysis_run_id": "run-001", "analysis_skill_snapshot_id": "triage-v1",
                "analysis_skill_execution_hash": "triage-exec-v1", "analysis_skill_version": "1",
                "analysis_output_sha256": "result-001", "analysis_model": "sonnet", "analysis_effort": "low",
                "report_fragment": {
                    "contract": "report-fragment-v1", "headline": "Payment timeout",
                    "severity": "high", "summary": {"zh": "支付超时", "en": "Payment timeout"},
                    "likely_cause": {"zh": "下游缓慢", "en": "Slow dependency"},
                    "recommended_action": {"zh": "检查依赖", "en": "Inspect dependency"},
                    "findings": [],
                },
            },
            {
                "id": "inc-002", "display_id": "INC-002", "title": "Cache errors",
                "status": "ACTIVE", "service": "cache", "environment": "prod",
                "triggered_at": "2026-09-14T02:00:00+00:00", "recovered_at": None,
                "trigger_value": None, "teams_url": None, "grafana_url": None, "log_filename": None,
                "screenshots": [], "analysis_run_id": "run-002", "analysis_skill_snapshot_id": "triage-v2",
                "analysis_skill_execution_hash": "triage-exec-v2", "analysis_skill_version": "2",
                "analysis_output_sha256": "result-002", "analysis_model": "sonnet", "analysis_effort": "medium",
                "report_fragment": {
                    "contract": "report-fragment-v1", "headline": "Cache errors",
                    "severity": "medium", "summary": {"zh": "缓存错误", "en": "Cache errors"},
                    "likely_cause": {"zh": "", "en": ""}, "recommended_action": {"zh": "", "en": ""},
                    "findings": [],
                },
            },
        ],
    }


def test_plan_resolves_frozen_facts_and_embeds_only_trusted_screenshot(tmp_path):
    plan = {
        "title": "Executive Overview",
        "blocks": [
            {"type": "heading", "text": "Critical Incidents", "level": 1},
            {"type": "incident_reference", "incident_id": "inc-001"},
            {"type": "analysis_reference", "analysis_run_id": "run-001"},
            {"type": "divider"},
            {"type": "bilingual_generated_text", "zh": "班次稳定。", "en": "The shift was stable."},
            {"type": "page_break"},
        ],
    }
    document = compose_report(plan, _snapshot())

    from noc_bridge.report_document import AnalysisReference, Heading, IncidentEvidence, PageBreak

    section_headings = [b.text for b in document.blocks if isinstance(b, Heading) and b.level == 1]
    assert [h for h in section_headings if h in ("Alerts", "General Summary", "Log Analysis")] == [
        "Alerts", "General Summary", "Log Analysis"
    ]

    incident_block = next(b for b in document.blocks if isinstance(b, IncidentEvidence))
    assert incident_block.incident_id == "inc-001"
    assert {link.label for link in incident_block.links} == {"Teams", "Grafana"}
    summary = next(b for b in document.blocks if isinstance(b, BilingualText))
    assert (summary.heading_zh, summary.heading_en) == ("Chinese Summary", "English Summary")

    analysis = next(b for b in document.blocks if isinstance(b, AnalysisReference))
    assert analysis.analysis_run_id == "run-001"
    assert any(item.label == "Execution Hash" for item in analysis.provenance)

    # Canonical section order: Alerts -> General Summary -> Log Analysis,
    # each separated by an explicit page break, regardless of how the plan
    # interleaved its blocks.
    alerts_idx = document.blocks.index(incident_block)
    analysis_idx = document.blocks.index(analysis)
    assert alerts_idx < analysis_idx
    page_breaks = [i for i, b in enumerate(document.blocks) if isinstance(b, PageBreak)]
    assert len(page_breaks) == 2
    assert alerts_idx < page_breaks[0] < analysis_idx
    assert page_breaks[0] < page_breaks[1] < analysis_idx

    destination = tmp_path / "report.docx"
    render_document(document, destination, screenshot_fetcher=lambda bucket, key: _PNG)
    text = "\n".join(paragraph.text for paragraph in Document(str(destination)).paragraphs)
    assert "Trigger Value: 95%" in text
    assert "Payment timeout" in text
    with zipfile.ZipFile(destination) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())


def test_unknown_references_are_rejected():
    with pytest.raises(OutputValidationError, match="unknown incident reference"):
        compose_report({"blocks": [{"type": "incident_reference", "incident_id": "secret"}]}, _snapshot())
    with pytest.raises(OutputValidationError, match="outside this report"):
        compose_report({"blocks": [{"type": "analysis_reference", "analysis_run_id": "run-999"}]}, _snapshot())


def test_report_composition_enforces_semantic_coverage_and_duplicates():
    snapshot = _snapshot()
    snapshot["coverage"] = {"incidents": "all", "analyses": "all_available"}
    incomplete = {"blocks": [
        {"type": "incident_reference", "incident_id": "inc-001"},
        {"type": "analysis_reference", "analysis_run_id": "run-001"},
    ]}
    with pytest.raises(OutputValidationError, match="omitted required incidents"):
        compose_report(incomplete, snapshot)

    complete = {"blocks": [
        {"type": "incident_reference", "incident_id": "inc-001"},
        {"type": "incident_reference", "incident_id": "inc-002"},
        {"type": "analysis_reference", "analysis_run_id": "run-001"},
        {"type": "analysis_reference", "analysis_run_id": "run-002"},
    ]}
    compose_report(complete, snapshot)
    duplicate = {"blocks": complete["blocks"] + [{"type": "incident_reference", "incident_id": "inc-001"}]}
    with pytest.raises(OutputValidationError, match="duplicate incident"):
        compose_report(duplicate, snapshot)


def test_full_analysis_presentation_keeps_all_findings_without_ai_echo(tmp_path):
    snapshot = _snapshot()
    snapshot["incidents"][0]["analysis"] = {
        "summary_zh": "完整摘要", "summary_en": "Full summary",
        "key_finds": [
            {"label_zh": f"重点{i}", "label_en": f"Finding {i}", "count": i, "percentage": i * 10, "detail_zh": f"细节{i}", "detail_en": f"Detail {i}"}
            for i in range(1, 8)
        ],
        "secondary_finds": [
            {"label_zh": f"次要{i}", "label_en": f"Secondary {i}", "count": i, "percentage": i * 5, "detail_zh": f"次要细节{i}", "detail_en": f"Secondary detail {i}"}
            for i in range(1, 5)
        ],
        "likely_cause_zh": "根因", "likely_cause_en": "Root cause",
        "recommended_action_zh": "修复", "recommended_action_en": "Remediate",
    }
    document = compose_report(
        {"blocks": [
            {"type": "incident_reference", "incident_id": "inc-001"},
            {"type": "analysis_reference", "analysis_run_id": "run-001"},
        ]},
        snapshot,
    )
    destination = tmp_path / "phase7-analysis-fidelity.docx"
    render_document(document, destination, screenshot_fetcher=lambda *_: _PNG)
    text = "\n".join(paragraph.text for paragraph in Document(str(destination)).paragraphs)
    for value in ("Finding 1", "Finding 7", "Secondary 1", "Secondary 4", "Root cause", "Remediate", "7 / 70%"):
        assert value in text

    from noc_bridge.report_document import FindList
    analysis_block = next(block for block in document.blocks if block.__class__.__name__ == "AnalysisReference")
    children = list(analysis_block.children)
    headings = [child.text for child in children if hasattr(child, "text")]
    assert headings.index("Chinese") < headings.index("English")
    chinese_end = headings.index("English")
    assert all(getattr(child, "language", None) == "Chinese" for child in children[:chinese_end] if isinstance(child, FindList))
    assert all(getattr(child, "language", None) == "English" for child in children[chinese_end:] if isinstance(child, FindList))


def test_canonical_profile_numbers_from_frozen_order_and_repeats_alert_evidence(tmp_path):
    snapshot = _snapshot()
    snapshot["composition_profile"] = "noc-daily-report-v1"
    snapshot["coverage"] = {
        "incidents": "all", "analyses": "all_available",
        "allow_duplicate_incidents": False, "allow_duplicate_analyses": False,
    }
    # Claude returns narrative only; coverage and order come from the frozen
    # snapshot and profile.
    plan = {"general_summary": {"zh": "班次摘要", "en": "Shift summary"}}
    document = compose_report(plan, snapshot)
    from noc_bridge.report_document import AnalysisReference, IncidentEvidence

    assert document.title == "Daily Alert & Log Analysis Report"
    assert next(item.value for item in document.metadata if item.label == "Shift") == "MS"
    alerts = [b for b in document.blocks if isinstance(b, IncidentEvidence)]
    analyses = [b for b in document.blocks if isinstance(b, AnalysisReference)]
    assert [b.heading for b in alerts] == ["Alert #1 - Payment timeout", "Alert #2 - Cache errors"]
    assert [b.heading for b in analyses] == ["Alert #1 - Payment timeout", "Alert #2 - Cache errors"]
    assert len(alerts[0].screenshots) == 1
    assert len(analyses[0].screenshots) == 1
    assert alerts[0].metadata == ()
    assert {link.label for link in alerts[0].links} == {"Grafana"}
    destination = tmp_path / "canonical.docx"
    render_document(document, destination, screenshot_fetcher=lambda *_: _PNG)
    rendered_text = "\n".join(paragraph.text for paragraph in Document(str(destination)).paragraphs)
    assert "Execution Hash" not in rendered_text
    with zipfile.ZipFile(destination) as archive:
        relationships = archive.read("word/_rels/document.xml.rels").decode()
        assert "https://grafana.example/inc-001" in relationships
        assert "TargetMode=\"External\"" in relationships


def test_narrative_plan_derives_complete_canonical_coverage_without_references():
    snapshot = _snapshot()
    snapshot["composition_profile"] = "noc-daily-report-v1"
    snapshot["coverage"] = {"incidents": "all", "analyses": "all_available"}
    document = compose_report(
        {"general_summary": {"zh": "班次稳定。", "en": "The shift was stable."}},
        snapshot,
    )
    from noc_bridge.report_document import AnalysisReference, IncidentEvidence
    assert [b.incident_id for b in document.blocks if isinstance(b, IncidentEvidence)] == ["inc-001", "inc-002"]
    assert [b.incident_id for b in document.blocks if isinstance(b, AnalysisReference)] == ["inc-001", "inc-002"]


@pytest.mark.parametrize("plan", [
    {},
    {"general_summary": {"zh": "", "en": "English"}},
    {"general_summary": {"zh": "中文", "en": "   "}},
])
def test_canonical_daily_report_requires_nonempty_bilingual_general_summary(plan):
    snapshot = _snapshot()
    snapshot["composition_profile"] = "noc-daily-report-v1"
    snapshot["coverage"] = {"incidents": "all", "analyses": "all_available"}
    with pytest.raises(OutputValidationError, match="general_summary"):
        compose_report(plan, snapshot)


def test_schema_rejects_a_plan_without_the_canonical_summary():
    with pytest.raises(OutputValidationError, match="required property"):
        validate_output("daily_report", {}, skill_name="daily-alert-report", skills_dir=SKILLS_DIR)


def test_materially_different_analysis_schema_uses_generic_presentation_export(tmp_path):
    snapshot = _snapshot()
    snapshot["incidents"][0]["analysis"] = {
        "different_private_schema": {"primary_failure": "database"},
        "report_presentation": {
            "blocks": [
                {"type": "heading", "text": "Primary Failure", "level": 3},
                {"type": "bilingual_text", "zh": "数据库故障", "en": "Database failure"},
            ]
        },
    }
    snapshot["incidents"][0]["analysis_presentation_contract"] = "analysis-presentation-v1"
    document = compose_report(
        {"blocks": [{"type": "analysis_reference", "analysis_run_id": "run-001"}]},
        snapshot,
    )
    destination = tmp_path / "different-analysis-schema.docx"
    render_document(document, destination, screenshot_fetcher=lambda *_: _PNG)
    text = "\n".join(paragraph.text for paragraph in Document(str(destination)).paragraphs)
    assert "Primary Failure" in text
    assert "Database failure" in text


def test_storage_references_are_not_part_of_the_report_plan_contract():
    malicious = {
        "blocks": [{"type": "screenshot", "bucket": "other-bucket", "object_key": "secret/file"}]
    }
    with pytest.raises(Exception):
        validate_output("daily_report", malicious, skill_name="daily-alert-report", skills_dir=SKILLS_DIR)


def test_expected_screenshot_retrieval_failure_is_not_silently_successful(tmp_path):
    document = compose_report(
        {"blocks": [{"type": "incident_reference", "incident_id": "inc-001"}]},
        _snapshot(),
    )
    with pytest.raises(EvidenceRetrievalError):
        render_document(document, tmp_path / "temporary-failure.docx", screenshot_fetcher=lambda *_: None)
    with pytest.raises(EvidenceIntegrityError):
        render_document(document, tmp_path / "missing-fetcher.docx")


def test_no_frozen_screenshot_is_a_valid_report(tmp_path):
    snapshot = _snapshot()
    snapshot["incidents"][0]["screenshots"] = []
    document = compose_report(
        {"blocks": [{"type": "incident_reference", "incident_id": "inc-001"}]},
        snapshot,
    )
    render_document(document, tmp_path / "no-screenshot.docx")


def _many_incidents(n: int) -> dict:
    snapshot = {
        "shift_starts_at": "2026-09-13T00:00:00+00:00",
        "shift_ends_at": "2026-09-13T08:00:00+00:00",
        "report_skill_snapshot_id": "report-vN",
        "report_skill_execution_hash": "report-exec-vN",
        "incidents": [],
    }
    for i in range(1, n + 1):
        snapshot["incidents"].append({
            "id": f"inc-{i:03d}", "display_id": f"INC-{i:03d}", "title": f"Service {i} failure",
            "status": "RECOVERED", "service": f"service-{i}", "environment": "prod",
            "triggered_at": "2026-09-13T01:00:00+00:00", "recovered_at": "2026-09-13T01:30:00+00:00",
            "trigger_value": None, "teams_url": None,
            "grafana_url": f"https://grafana.example/inc-{i:03d}",
            "log_filename": f"{i}-service-{i}-logs-2026-09-13.json",
            "screenshots": [], "analysis_run_id": f"run-{i:03d}",
            "analysis_skill_snapshot_id": "triage-v1", "analysis_skill_execution_hash": "triage-exec-v1",
            "analysis_skill_version": "1", "analysis_output_sha256": f"result-{i:03d}",
            "analysis_model": "sonnet", "analysis_effort": "low",
            "report_fragment": {
                "contract": "report-fragment-v1", "headline": f"Service {i} failure",
                "severity": "high", "summary": {"zh": f"服务{i}故障", "en": f"Service {i} failure"},
                "likely_cause": {"zh": "", "en": ""}, "recommended_action": {"zh": "", "en": ""},
                "findings": [],
            },
        })
    return snapshot


def test_golden_report_section_order_and_pagination_with_many_alerts(tmp_path):
    """Regression fixture for the canonical Daily Report layout: N alerts,
    each exactly once, Alerts -> General Summary -> Log Analysis, with no
    cap on how many alerts may appear (natural, unlimited pagination)."""
    from noc_bridge.report_document import AnalysisReference, Heading, IncidentEvidence, PageBreak

    n = 12  # comfortably more than would fit on a single page
    snapshot = _many_incidents(n)
    blocks = [{"type": "incident_reference", "incident_id": inc["id"]} for inc in snapshot["incidents"]]
    blocks.append({"type": "bilingual_generated_text", "zh": "班次总体稳定。", "en": "The shift was overall stable."})
    blocks.extend({"type": "analysis_reference", "incident_id": inc["id"]} for inc in snapshot["incidents"])

    document = compose_report({"blocks": blocks}, snapshot)

    incident_blocks = [b for b in document.blocks if isinstance(b, IncidentEvidence)]
    analysis_blocks = [b for b in document.blocks if isinstance(b, AnalysisReference)]

    # every alert appears, exactly once, in both Alerts and Log Analysis
    assert [b.incident_id for b in incident_blocks] == [inc["id"] for inc in snapshot["incidents"]]
    assert len(incident_blocks) == n
    assert len(analysis_blocks) == n

    section_headings = [b.text for b in document.blocks if isinstance(b, Heading) and b.level == 1]
    assert [h for h in section_headings if h in ("Alerts", "General Summary", "Log Analysis")] == [
        "Alerts", "General Summary", "Log Analysis"
    ]

    alerts_start = document.blocks.index(incident_blocks[0])
    alerts_end = document.blocks.index(incident_blocks[-1])
    analysis_start = document.blocks.index(analysis_blocks[0])

    # all N alerts land contiguously in the Alerts section (natural
    # pagination is a DOCX/word-wrap concern, not a block-ordering cap)
    assert alerts_end - alerts_start == n - 1
    # General Summary and its page break sit strictly between Alerts and
    # Log Analysis -- never interleaved with either.
    assert alerts_end < analysis_start
    page_breaks = [i for i, b in enumerate(document.blocks) if isinstance(b, PageBreak)]
    assert len(page_breaks) == 2
    assert alerts_end < page_breaks[0] < analysis_start
    assert page_breaks[0] < page_breaks[1] < analysis_start

    destination = tmp_path / "many-alerts.docx"
    render_document(document, destination)
    with zipfile.ZipFile(destination) as archive:
        xml = archive.read("word/document.xml").decode()
        # Only the two deterministic section boundaries are explicit. Alerts
        # are allowed to flow naturally across as many Word pages as needed.
        assert xml.count('w:type="page"') == 2
    text = "\n".join(paragraph.text for paragraph in Document(str(destination)).paragraphs)
    assert "Alert #12 - Service 12 failure" in text
