"""Skill Runtime mission Phase 8/9: ReportDocument assembly + generic
DOCX rendering. Proves the two things the mission actually cares about:
(1) build_daily_report_document translates a daily-alert-report result
into typed blocks with no surprises, and (2) render_document renders any
ReportDocument built purely from those block types, with zero knowledge
of where the blocks came from — the "format change flexibility" proof at
the ReportDocument layer.
"""
from __future__ import annotations

import pathlib
import base64

from docx import Document

from noc_bridge.docx_render import render_daily_report_docx, render_document
from noc_bridge.report_composition import compose_report
from noc_bridge.report_document import (
    AnalysisReference,
    BilingualFindList,
    BilingualText,
    Heading,
    FindList,
    Find,
    IncidentEvidence,
    Metadata,
    Paragraph,
    ReportDocument,
    build_report_document,
    build_daily_report_document,
    DOCUMENT_BLOCK_TYPES,
)
from noc_bridge.validation import validate_output

SKILLS_DIR = pathlib.Path(__file__).resolve().parents[2] / "skills"
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)

_RESULT = {
    "title": "Daily Alert Report — 2026-09-13",
    "overview_en": "Two incidents this shift.",
    "overview_zh": "本班次两起事件。",
    "cross_incident_findings_en": "No correlation found.",
    "cross_incident_findings_zh": "未发现关联。",
    "sections": [
        {
            "incident_display_id": "INC-001",
            "title": "Payment service errors",
            "status": "RESOLVED",
            "grafana_url": "https://grafana.example/d/abc",
            "log_filename": "payments.log",
            "screenshots": [{"bucket": "noc-evidence", "object_key": "k1", "filename": "s1.png"}],
            "analysis": {
                "summary_en": "NullPointerException in PaymentWorker",
                "summary_zh": "支付服务空指针异常",
                "key_finds": [
                    {"label_en": "NPE", "label_zh": "空指针", "detail_en": "d", "detail_zh": "d",
                     "count": 5, "percentage": 100.0}
                ],
                "secondary_finds": [],
                "likely_cause_en": "null pointer", "likely_cause_zh": "空指针",
                "recommended_action_en": "patch", "recommended_action_zh": "修复",
            },
        },
        {
            "incident_display_id": "INC-002",
            "title": "Unanalyzed incident",
            "status": "OPEN",
            "grafana_url": None,
            "log_filename": None,
            "screenshots": [],
            "analysis": None,
        },
    ],
}


def test_build_daily_report_document_produces_expected_block_shape():
    document = build_daily_report_document(_RESULT)
    assert isinstance(document, ReportDocument)
    assert document.title == _RESULT["title"]

    headings = [b.text for b in document.blocks if isinstance(b, Heading)]
    assert headings == ["Alerts", "General Summary", "Cross-Incident Findings", "Log Analysis"]

    evidence_blocks = [b for b in document.blocks if isinstance(b, IncidentEvidence)]
    assert len(evidence_blocks) == 2
    assert evidence_blocks[0].metadata == (Metadata("Incident", "INC-001"), Metadata("Status", "RESOLVED"))
    assert evidence_blocks[0].link.url == "https://grafana.example/d/abc"
    assert evidence_blocks[1].link is None  # no grafana_url for INC-002

    analysis_blocks = [b for b in document.blocks if isinstance(b, AnalysisReference)]
    assert len(analysis_blocks) == 2
    assert analysis_blocks[0].available is True
    assert analysis_blocks[0].key_finds.finds_en[0].label == "NPE"
    assert analysis_blocks[1].available is False


def test_build_daily_report_document_handles_no_incidents():
    result = {**_RESULT, "sections": []}
    document = build_daily_report_document(result)
    paragraphs = [b for b in document.blocks if isinstance(b, Paragraph)]
    assert any(p.text == "No incidents in this shift." for p in paragraphs)


def test_render_document_is_generic_over_a_synthetic_document_shape(tmp_path):
    """Proves render_document has no dependency on daily-alert-report's
    own assembly function — a hand-built ReportDocument using a subset of
    block types (unrelated to any real skill) still renders to a valid
    DOCX. This is the format-change-flexibility proof at the renderer
    layer: a wholly different report shape needs no renderer changes."""
    document = ReportDocument(
        title="Synthetic Report",
        blocks=(
            Heading("Section One", level=1),
            Paragraph("Just some text."),
            BilingualText(text_zh="你好", text_en="Hello", heading_zh="ZH", heading_en="EN"),
        ),
    )
    dest = tmp_path / "synthetic.docx"
    render_document(document, dest)
    assert dest.exists()

    doc = Document(str(dest))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Just some text." in all_text
    assert "你好" in all_text
    assert "Hello" in all_text


def test_every_documented_block_type_is_parsed_and_rendered(tmp_path):
    raw_blocks = [
        {"type": "heading", "text": "H", "level": 1},
        {"type": "paragraph", "text": "P"},
        {"type": "divider"},
        {"type": "page_break"},
        {"type": "metadata", "label": "M", "value": "V"},
        {"type": "link", "label": "Grafana", "url": "https://example.test", "text": "Grafana"},
        {"type": "log_file_reference", "filename": "log.json"},
        {"type": "screenshot", "bucket": "b", "object_key": "k", "filename": "s.png"},
        {"type": "bilingual_text", "zh": "中", "en": "En"},
        {"type": "bilingual_find_list", "heading_zh": "重点", "heading_en": "Findings", "finds_zh": [], "finds_en": []},
        {"type": "find_list", "heading": "Key Finds", "language": "Chinese", "finds": [{"label": "x", "detail": "d", "stat": "1 / 2%"}]},
        {"type": "incident_evidence", "heading": "Alert", "metadata": [], "links": [], "screenshots": []},
        {"type": "analysis_reference", "heading": "Analysis", "available": True, "children": []},
    ]
    assert {block["type"] for block in raw_blocks} == set(DOCUMENT_BLOCK_TYPES)
    document = build_report_document({"metadata": {"title": "Contract"}, "blocks": raw_blocks})
    assert len(document.blocks) == len(raw_blocks)
    assert isinstance(document.blocks[9], BilingualFindList)
    assert isinstance(document.blocks[10], FindList)
    destination = tmp_path / "contract.docx"
    render_document(document, destination, screenshot_fetcher=lambda *_: _PNG)
    assert destination.exists()


def test_materially_different_skill_layouts_use_same_docx_adapter(tmp_path):
    """Acceptance proof: section order and meaning come from each skill's
    declarative block result; render_document remains unchanged."""
    layouts = [
        {
            "metadata": {"title": "Layout A"},
            "blocks": [
                {"type": "heading", "level": 1, "text": "Alerts"},
                {"type": "paragraph", "text": "General Summary"},
                {"type": "heading", "level": 1, "text": "Log Analysis"},
            ],
        },
        {
            "metadata": {"title": "Layout B"},
            "blocks": [
                {"type": "heading", "level": 1, "text": "Executive Overview"},
                {"type": "divider"},
                {"type": "heading", "level": 1, "text": "Critical Incidents"},
                {"type": "paragraph", "text": "Recurring Problems"},
                {"type": "page_break"},
                {"type": "heading", "level": 1, "text": "Recommended Actions"},
            ],
        },
    ]
    for index, result in enumerate(layouts):
        dest = tmp_path / f"layout-{index}.docx"
        render_document(build_report_document(result), dest)
        text = "\n".join(p.text for p in Document(str(dest)).paragraphs)
        assert result["metadata"]["title"] in text
        assert result["blocks"][0]["text"] in text


def test_declarative_document_supports_metadata_and_analysis_references(tmp_path):
    result = {
        "metadata": {"title": "Executive Report", "date": "2026-09-14", "shift": "Night"},
        "blocks": [
            {"type": "heading", "level": 1, "text": "Critical Incidents"},
            {"type": "analysis_reference", "analysis_run_id": "run-307", "label": "Incident 1"},
        ],
    }
    destination = tmp_path / "provenance.docx"
    render_document(build_report_document(result), destination)

    text = "\n".join(p.text for p in Document(str(destination)).paragraphs)
    assert "date: 2026-09-14" in text
    assert "Incident 1" in text
    assert "Analysis reference: run-307" not in text

    audit_destination = tmp_path / "provenance-audit.docx"
    render_document(build_report_document(result), audit_destination, include_provenance=True)
    audit_text = "\n".join(p.text for p in Document(str(audit_destination)).paragraphs)
    assert "Analysis reference: run-307" in audit_text


def test_active_daily_report_document_contract_renders_nested_analysis_blocks(tmp_path):
    result = {
        "general_summary": {"zh": "一项事件。", "en": "One incident."},
    }
    validate_output(
        "daily_report",
        result,
        skill_name="daily-alert-report",
        skills_dir=SKILLS_DIR,
    )
    snapshot = {
        "shift_starts_at": "2026-09-14T00:00:00+00:00",
        "shift_ends_at": "2026-09-14T08:00:00+00:00",
        "shift_timezone": "Asia/Manila",
        "report_skill_snapshot_id": "report-snapshot",
        "report_skill_execution_hash": "report-execution",
        "incidents": [{
            "id": "incident-1", "display_id": "INC-001", "title": "API failure",
            "status": "RESOLVED", "service": "api", "environment": "prod",
            "triggered_at": "2026-09-14T01:00:00+00:00", "recovered_at": None,
            "grafana_url": "https://grafana.example/1", "log_filename": "api.log",
            "screenshots": [], "analysis_run_id": "run-1",
            "report_fragment": {"summary": {"zh": "数据库超时。", "en": "Database timeout."}, "findings": [], "likely_cause": {}, "recommended_action": {}},
            "analysis_skill_snapshot_id": "analysis-snapshot", "analysis_skill_execution_hash": "analysis-execution",
            "analysis_skill_version": "1", "analysis_output_sha256": "result-hash", "analysis_model": "model", "analysis_effort": "low",
        }],
        "composition_profile": "noc-daily-report-v1",
        "coverage": {"incidents": "all", "analyses": "all_available"},
    }
    destination = tmp_path / "active-daily-report.docx"
    render_document(compose_report(result, snapshot), destination)
    text = "\n".join(p.text for p in Document(str(destination)).paragraphs)
    assert "Incident ID: incident-1" not in text
    assert "Analysis reference: run-1" not in text
    assert "Database timeout." in text


def test_render_daily_report_docx_end_to_end_produces_a_valid_docx(tmp_path):
    dest = tmp_path / "report.docx"
    fetched = []

    def fetcher(bucket, object_key):
        fetched.append((bucket, object_key))
        return _PNG

    render_daily_report_docx(_RESULT, dest, screenshot_fetcher=fetcher)
    assert dest.exists()
    # The canonical Daily Report repeats trusted screenshots in Alerts and
    # Log Analysis; both renderings must resolve the frozen object.
    assert fetched == [("noc-evidence", "k1"), ("noc-evidence", "k1")]

    doc = Document(str(dest))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "INC-001" in all_text
    assert "NullPointerException in PaymentWorker" in all_text
    assert "No log analysis available yet for this incident." in all_text
