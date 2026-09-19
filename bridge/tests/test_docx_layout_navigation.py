from __future__ import annotations

import io
import zipfile

from docx import Document
from PIL import Image

from noc_bridge.docx_render import (
    _PAIR_MAX_HEIGHT_IN,
    _PAIR_MAX_WIDTH_IN,
    _SINGLE_MAX_HEIGHT_IN,
    _SINGLE_MAX_WIDTH_IN,
    render_document,
)
from noc_bridge.report_composition import compose_report
from noc_bridge.report_document import AnalysisReference, Heading, IncidentEvidence, ReportDocument, Screenshot


EMU_PER_INCH = 914400


def _png(width: int, height: int, *, dpi: int = 100) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(output, format="PNG", dpi=(dpi, dpi))
    return output.getvalue()


def _extents(path) -> list[tuple[int, int]]:
    doc = Document(str(path))
    return [
        (int(node.get("cx")), int(node.get("cy")))
        for node in doc.element.xpath(".//wp:extent")
    ]


def _evidence(*shots: Screenshot) -> ReportDocument:
    return ReportDocument("Layout", (IncidentEvidence("Alert #1 - Example", screenshots=shots),))


def test_single_screenshot_is_bounded_aspect_correct_centered_and_not_upscaled(tmp_path):
    large = _png(1000, 800)
    path = tmp_path / "large.docx"
    render_document(_evidence(Screenshot("b", "large", "large.png")), path, lambda *_: large)
    width, height = _extents(path)[0]
    assert width <= _SINGLE_MAX_WIDTH_IN * EMU_PER_INCH
    assert height <= _SINGLE_MAX_HEIGHT_IN * EMU_PER_INCH
    assert abs((width / height) - (1000 / 800)) < 0.01
    drawing_paragraph = next(p for p in Document(str(path)).paragraphs if p._p.xpath(".//w:drawing"))
    assert drawing_paragraph.alignment == 1

    small = _png(100, 50)
    path = tmp_path / "small.docx"
    render_document(_evidence(Screenshot("b", "small", "small.png")), path, lambda *_: small)
    width, height = _extents(path)[0]
    assert width <= 1.01 * EMU_PER_INCH
    assert height <= 0.51 * EMU_PER_INCH


def test_trigger_recovery_pair_uses_compact_columns_and_portraits_stack(tmp_path):
    landscape = {"trigger": _png(900, 500), "recover": _png(900, 500)}
    pair = _evidence(
        Screenshot("b", "trigger", "Triggered.png"),
        Screenshot("b", "recover", "Recovered.png"),
    )
    path = tmp_path / "pair.docx"
    render_document(pair, path, lambda _b, key, *_: landscape[key])
    doc = Document(str(path))
    assert len(doc.tables) == 1
    assert [cell.text for cell in doc.tables[0].rows[0].cells] == ["Triggered", "Recovered"]
    assert all(
        width <= _PAIR_MAX_WIDTH_IN * EMU_PER_INCH and height <= _PAIR_MAX_HEIGHT_IN * EMU_PER_INCH
        for width, height in _extents(path)
    )

    portrait_data = {"trigger": _png(500, 1200), "recover": _png(500, 1200)}
    path = tmp_path / "portrait.docx"
    render_document(pair, path, lambda _b, key, *_: portrait_data[key])
    assert len(Document(str(path)).tables) == 0
    assert len(_extents(path)) == 2

    analysis_path = tmp_path / "pair-analysis.docx"
    analysis_doc = ReportDocument(
        "Layout",
        (
            Heading("Log Analysis", level=1),
            AnalysisReference(
                "Alert #1 - Example",
                available=True,
                screenshots=pair.blocks[0].screenshots,
                children=(),
            ),
        ),
    )
    render_document(analysis_doc, analysis_path, lambda _b, key, *_: landscape[key])
    assert len(Document(str(analysis_path)).tables) == 1


def _snapshot(count: int = 3) -> dict:
    incidents = []
    for index in range(1, count + 1):
        incidents.append({
            "id": f"incident-id-{index}",
            "display_id": f"INC-{index:03d}",
            "title": "重复告警！" if index <= 2 else f"Alert, punctuation #{index}",
            "status": "RECOVERED",
            "service": "api",
            "environment": "prod",
            "triggered_at": "2026-09-19T00:00:00+00:00",
            "recovered_at": "2026-09-19T00:05:00+00:00",
            "grafana_url": None,
            "log_filename": f"alert-{index}.json",
            "screenshots": [],
            "analysis_run_id": f"run-{index}",
            "report_fragment": {
                "summary": {"zh": "摘要", "en": "Summary"},
                "findings": [], "likely_cause": {}, "recommended_action": {},
            },
        })
    return {
        "shift_starts_at": "2026-09-19T00:00:00+00:00",
        "shift_ends_at": "2026-09-19T08:00:00+00:00",
        "shift_timezone": "Asia/Manila",
        "incidents": incidents,
        "composition_profile": "noc-daily-report-v1",
        "coverage": {"incidents": "all", "analyses": "all_available"},
    }


def test_navigation_targets_match_unique_bookmarks_for_duplicate_titles(tmp_path):
    document = compose_report({"general_summary": {"zh": "摘要", "en": "Summary"}}, _snapshot())
    path = tmp_path / "navigation.docx"
    render_document(document, path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode()
    doc = Document(str(path))
    names = [node.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name")
             for node in doc.element.xpath(".//w:bookmarkStart")]
    bookmark_ids = [node.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id")
                    for node in doc.element.xpath(".//w:bookmarkStart")]
    analysis_names = [name for name in names if name.startswith("log_analysis_")]
    anchors = [node.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor")
               for node in doc.element.xpath(".//w:hyperlink[@w:anchor]")]
    assert len(analysis_names) == 3
    assert len(set(analysis_names)) == 3
    assert len(bookmark_ids) == len(set(bookmark_ids))
    assert all(len(name) <= 40 and name[0].isalpha() for name in analysis_names)
    assert all(name in anchors for name in analysis_names)
    assert anchors.count(analysis_names[0]) == 2  # navigation + evidence heading
    assert anchors.count(analysis_names[1]) == 2
    assert "alerts_start" in names
    assert anchors.count("alerts_start") == 3  # one Back to Alerts link per analysis
    assert "Cross-Incident Findings" not in xml


def test_alert_without_usable_analysis_has_no_dead_navigation_link(tmp_path):
    snapshot = _snapshot(2)
    snapshot["incidents"][1]["analysis_run_id"] = None
    document = compose_report({"general_summary": {"zh": "摘要", "en": "Summary"}}, snapshot)
    path = tmp_path / "partial.docx"
    render_document(document, path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode()
    assert "Alert #1 - 重复告警！" in "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert xml.count('w:anchor="log_analysis_') == 2


def test_canonical_report_omits_empty_navigation_when_no_analysis_exists(tmp_path):
    snapshot = _snapshot(1)
    snapshot["incidents"][0]["analysis_run_id"] = None
    document = compose_report({"general_summary": {"zh": "摘要", "en": "Summary"}}, snapshot)
    path = tmp_path / "no-analysis.docx"
    render_document(document, path)
    text = "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
    assert "Alert Navigation" not in text
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode()
    assert 'w:anchor="log_analysis_' not in xml


def test_twelve_alert_navigation_targets_remain_unique(tmp_path):
    document = compose_report({"general_summary": {"zh": "摘要", "en": "Summary"}}, _snapshot(12))
    path = tmp_path / "twelve.docx"
    render_document(document, path)
    doc = Document(str(path))
    names = [node.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name")
             for node in doc.element.xpath(".//w:bookmarkStart")]
    analysis_names = [name for name in names if name.startswith("log_analysis_")]
    assert len(analysis_names) == len(set(analysis_names)) == 12
