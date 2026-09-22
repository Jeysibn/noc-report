"""Renders a `ReportDocument` (see noc_bridge.report_document) into a DOCX
file (master plan §29 Daily Report Job Contract: "... -> daily-alert-
report -> DOCX -> MinIO -> UI download"). This runs in the bridge process
itself, not inside the sandbox — the sandbox's only job is to produce the
structured JSON via the Claude CLI; keeping DOCX rendering out of the
sandbox image means that image doesn't need python-docx or any of its
dependencies baked in, and a rendering bug can be fixed/redeployed
without touching the sandbox image at all.

Skill Runtime mission Phase 9: `render_document` knows only
`report_document`'s block types. Translation from a skill-owned result to
those blocks lives in a report assembler, so a different report skill/shape
only needs its own assembler; this DOCX adapter is unchanged. The legacy
daily-report wrapper remains as a compatibility entry point for that one
assembler.
"""
from __future__ import annotations

import io
import pathlib
from typing import Callable

from docx import Document
from docx.image.image import Image
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

from noc_bridge.report_document import (
    AnalysisReference,
    AlertNavigation,
    BilingualFindList,
    BilingualText,
    Divider,
    FindList,
    Heading,
    IncidentEvidence,
    Link,
    LogFileReference,
    Metadata,
    Paragraph,
    PageBreak,
    ReportDocument,
    Screenshot,
    build_daily_report_document,
)
from noc_bridge.failures import EvidenceIntegrityError, EvidenceRetrievalError

ScreenshotFetcher = Callable[..., bytes | None]
_EMU_PER_INCH = 914400
_SINGLE_MAX_WIDTH_IN = 4.9
_SINGLE_MAX_HEIGHT_IN = 3.4
_PAIR_MAX_WIDTH_IN = 3.0
_PAIR_MAX_HEIGHT_IN = 3.2
_REFERENCE_FONT = "Noto Sans CJK SC"
_BODY_COLOR = RGBColor(26, 26, 26)
_MUTED_COLOR = RGBColor(102, 102, 102)
_ALERT_COLOR = RGBColor(46, 90, 172)


def _set_east_asia_font(style, name: str = _REFERENCE_FONT) -> None:
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _configure_document(doc: Document, title: str, metadata: tuple[Metadata, ...]) -> None:
    """Apply the compact, reference-report typography and page geometry."""
    for section in doc.sections:
        # Match the reference report's usable portrait page area (roughly
        # 0.65in side margins and 0.62in top/bottom margins). Screenshot
        # bounds remain deliberately smaller than this area.
        section.top_margin = Inches(0.62)
        section.bottom_margin = Inches(0.62)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        header = section.header.paragraphs[0]
        header.text = title
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(100, 116, 139)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date = next((m.value for m in metadata if m.label == "Date"), "")
        shift = next((m.value for m in metadata if m.label == "Shift"), "")
        footer.add_run(" · ".join(value for value in (date, shift) if value))
        footer.add_run("  |  Page ")
        page_field = OxmlElement("w:fldSimple")
        page_field.set(qn("w:instr"), "PAGE")
        footer._p.append(page_field)
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(100, 116, 139)

    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Intense Quote"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        _set_east_asia_font(style)
        if style_name == "Normal":
            style.font.size = Pt(10)
            style.font.color.rgb = _BODY_COLOR
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.line_spacing = 1.08
        elif style_name == "Title":
            style.font.size = Pt(20)
            style.font.bold = True
            style.font.italic = True
            style.font.color.rgb = _BODY_COLOR
        elif style_name == "Heading 1":
            style.font.size = Pt(16)
            style.font.bold = True
            style.font.italic = True
            style.font.color.rgb = _BODY_COLOR
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True
        elif style_name == "Heading 2":
            style.font.size = Pt(13)
            style.font.bold = True
            style.font.italic = True
            style.font.color.rgb = _ALERT_COLOR
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(5)
            style.paragraph_format.keep_with_next = True
        elif style_name == "Heading 3":
            style.font.size = Pt(11.5)
            style.font.bold = True
            style.font.italic = True
            style.font.color.rgb = _BODY_COLOR
            style.paragraph_format.space_before = Pt(8)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.keep_with_next = True
        elif style_name == "Intense Quote":
            style.font.size = Pt(10.5)
            style.font.bold = True
            style.font.italic = True
            style.font.color.rgb = RGBColor(51, 51, 51)
            style.paragraph_format.space_before = Pt(6)
            style.paragraph_format.space_after = Pt(3)


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2563EB")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_internal_hyperlink(paragraph, target: str, text: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), target)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2563EB")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_bookmark(paragraph, name: str) -> None:
    # Bookmark IDs need only match within a pair. The name is already a
    # deterministic SHA-based identifier, so its digest also provides a stable
    # document-local numeric ID.
    bookmark_id = "0" if name == "alerts_start" else str(int(name.rsplit("_", 1)[-1], 16))
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_link_paragraph(doc: Document, link: Link) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(3)
    # The reference report presents compact, human-readable link labels
    # instead of printing long dashboard URLs into the evidence block.
    _add_hyperlink(paragraph, link.url, link.prefix or link.label)


def _fetch_screenshot(shot: Screenshot, screenshot_fetcher: ScreenshotFetcher | None) -> bytes:
    if screenshot_fetcher is None:
        raise EvidenceIntegrityError(f"no screenshot fetcher for frozen evidence: {shot.filename or shot.object_key}")
    try:
        data = screenshot_fetcher(shot.bucket, shot.object_key, shot.version_id, shot.sha256)
    except TypeError:
        # Compatibility for renderer tests and historical adapters that only
        # understood the pre-versioned two-argument seam.
        data = screenshot_fetcher(shot.bucket, shot.object_key)
    if not data:
        raise EvidenceRetrievalError(f"frozen screenshot could not be retrieved: {shot.bucket}/{shot.object_key}")
    return data


def _scaled_dimensions(data: bytes, max_width_in: float, max_height_in: float) -> tuple[int, int]:
    image = Image.from_blob(data)
    width_in = image.px_width / float(image.horz_dpi or 96)
    height_in = image.px_height / float(image.vert_dpi or 96)
    scale = min(max_width_in / width_in, max_height_in / height_in, 1.0)
    return round(width_in * scale * _EMU_PER_INCH), round(height_in * scale * _EMU_PER_INCH)


def _add_picture(paragraph, data: bytes, max_width_in: float, max_height_in: float) -> tuple[int, int]:
    width, height = _scaled_dimensions(data, max_width_in, max_height_in)
    paragraph.add_run().add_picture(io.BytesIO(data), width=width, height=height)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return width, height


def _add_picture_dimensions(paragraph, data: bytes, width: int, height: int) -> None:
    paragraph.add_run().add_picture(io.BytesIO(data), width=width, height=height)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_screenshot(doc: Document, shot: Screenshot, screenshot_fetcher: ScreenshotFetcher | None) -> None:
    data = _fetch_screenshot(shot, screenshot_fetcher)
    try:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.keep_with_next = True
        _add_picture(paragraph, data, _SINGLE_MAX_WIDTH_IN, _SINGLE_MAX_HEIGHT_IN)
    except Exception as exc:
        raise EvidenceIntegrityError(f"frozen screenshot is not a valid image: {shot.filename or shot.object_key}") from exc


def _screenshot_label(shot: Screenshot, index: int) -> str:
    name = (shot.filename or shot.object_key).casefold()
    if "recover" in name:
        return "Recovered"
    if "trigger" in name:
        return "Triggered"
    return f"Screenshot {index}"


def _add_screenshots(doc: Document, shots: tuple[Screenshot, ...], screenshot_fetcher: ScreenshotFetcher | None) -> None:
    if len(shots) != 2:
        for shot in shots:
            _add_screenshot(doc, shot, screenshot_fetcher)
        return
    data = [_fetch_screenshot(shot, screenshot_fetcher) for shot in shots]
    try:
        dimensions = [_scaled_dimensions(item, _PAIR_MAX_WIDTH_IN, _PAIR_MAX_HEIGHT_IN) for item in data]
        # Very narrow portrait images become illegible in two columns. Stack
        # those at the compact single-image bounds instead.
        if any(width / _EMU_PER_INCH < 2.15 for width, _ in dimensions):
            for index, (shot, item) in enumerate(zip(shots, data), start=1):
                label = doc.add_paragraph(_screenshot_label(shot, index))
                label.alignment = WD_ALIGN_PARAGRAPH.CENTER
                label.paragraph_format.keep_with_next = True
                label.paragraph_format.space_after = Pt(2)
                _add_picture(doc.add_paragraph(), item, _SINGLE_MAX_WIDTH_IN, _SINGLE_MAX_HEIGHT_IN)
            return
        table = doc.add_table(rows=2, cols=2)
        table.autofit = False
        common_height = min(height for _, height in dimensions)
        for index, (shot, item) in enumerate(zip(shots, data)):
            label = table.cell(0, index).paragraphs[0]
            label.add_run(_screenshot_label(shot, index + 1)).bold = True
            label.alignment = WD_ALIGN_PARAGRAPH.CENTER
            label.paragraph_format.space_after = Pt(2)
            width, height = dimensions[index]
            width = round(width * common_height / height)
            _add_picture_dimensions(table.cell(1, index).paragraphs[0], item, width, common_height)
    except Exception as exc:
        raise EvidenceIntegrityError("frozen screenshot pair contains an invalid image") from exc


def _add_alert_navigation(doc: Document, block: AlertNavigation) -> None:
    heading = doc.add_heading("Alert Navigation", level=3)
    heading.paragraph_format.keep_with_next = True
    _add_bookmark(heading, "alerts_start")
    for entry in block.entries:
        paragraph = doc.add_paragraph(style="List Bullet")
        _add_internal_hyperlink(paragraph, entry.target, entry.text)


def _add_log_file(doc: Document, log_file: LogFileReference) -> None:
    heading = doc.add_paragraph()
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(4)
    heading.paragraph_format.space_after = Pt(2)
    run = heading.add_run("Log File")
    run.bold = True
    run.italic = True
    run.font.size = Pt(10.5)
    filename = doc.add_paragraph()
    filename.paragraph_format.space_after = Pt(4)
    filename.add_run("File Name: ").bold = True
    if log_file.url:
        _add_hyperlink(filename, log_file.url, log_file.filename)
    else:
        filename.add_run(log_file.filename)


def _add_finds(doc: Document, heading: str, finds, *, zh: bool) -> None:
    if not finds:
        return
    heading_paragraph = doc.add_paragraph(heading, style="Intense Quote")
    heading_paragraph.paragraph_format.keep_with_next = True
    for i, find in enumerate(finds, start=1):
        stats = f"（{find.stat}）" if (zh and find.stat) else (f" ({find.stat})" if find.stat else "")
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.1)
        p.paragraph_format.right_indent = Inches(0.1)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(f"{i}. {find.label}{stats}").bold = True
        if find.detail:
            p.add_run(f" - {find.detail}")


def _add_bilingual_find_list(doc: Document, find_list: BilingualFindList) -> None:
    _add_finds(doc, find_list.heading_zh, find_list.finds_zh, zh=True)
    _add_finds(doc, find_list.heading_en, find_list.finds_en, zh=False)


def _add_bilingual_text(doc: Document, block: BilingualText) -> None:
    if block.heading_zh:
        doc.add_heading(block.heading_zh, level=3)
    zh = doc.add_paragraph(block.text_zh)
    zh.paragraph_format.left_indent = Inches(0.1)
    zh.paragraph_format.right_indent = Inches(0.1)
    if block.heading_en:
        doc.add_heading(block.heading_en, level=3)
    en = doc.add_paragraph(block.text_en)
    en.paragraph_format.left_indent = Inches(0.1)
    en.paragraph_format.right_indent = Inches(0.1)


def _add_incident_evidence(doc: Document, block: IncidentEvidence, screenshot_fetcher: ScreenshotFetcher | None, *, include_provenance: bool = False) -> None:
    # Keep the alert heading glued to whatever paragraph follows it so Word's
    # automatic pagination does not strand "Alert #n - <title>" alone at the
    # bottom of a page with its evidence pushed to the next one.
    heading = doc.add_heading("", level=2)
    heading.paragraph_format.keep_with_next = True
    if block.navigation_target:
        _add_internal_hyperlink(heading, block.navigation_target, block.heading)
    else:
        heading.add_run(block.heading)
    _add_screenshots(doc, block.screenshots, screenshot_fetcher)
    if block.metadata:
        for index, item in enumerate(block.metadata):
            metadata = doc.add_paragraph()
            metadata.paragraph_format.space_before = Pt(2 if index == 0 else 0)
            metadata.paragraph_format.space_after = Pt(2)
            metadata.add_run(f"{item.label}: ").bold = True
            metadata.add_run(item.value)
    links = list(block.links)
    if block.link:
        links.insert(0, block.link)
    for link in links:
        _add_link_paragraph(doc, link)
    if block.log_file:
        _add_log_file(doc, block.log_file)
    if include_provenance and block.incident_id:
        doc.add_paragraph(f"Incident ID: {block.incident_id}")


def _add_analysis_reference(
    doc: Document,
    block: AnalysisReference,
    screenshot_fetcher: ScreenshotFetcher | None,
    *,
    include_provenance: bool = False,
) -> None:
    heading = doc.add_heading(block.heading, level=2)
    heading.paragraph_format.keep_with_next = True
    if heading.runs:
        heading.runs[0].font.color.rgb = _ALERT_COLOR
    if block.bookmark:
        _add_bookmark(heading, block.bookmark)
    _add_screenshots(doc, block.screenshots, screenshot_fetcher)
    if block.log_file:
        _add_log_file(doc, block.log_file)
    if include_provenance:
        if block.analysis_run_id:
            doc.add_paragraph(f"Analysis reference: {block.analysis_run_id}")
        for metadata in (*block.metadata, *block.provenance):
            doc.add_paragraph(f"{metadata.label}: {metadata.value}")
    if not block.available:
        doc.add_paragraph(block.unavailable_text or "No analysis available.")
        return
    for child in block.children:
        _render_block(doc, child, screenshot_fetcher, include_provenance=include_provenance)
    # Compatibility fields remain readable for historical daily-report
    # snapshots created before analysis_reference.children existed.
    if block.summary:
        _add_bilingual_text(doc, block.summary)
    if block.key_finds:
        _add_bilingual_find_list(doc, block.key_finds)
    if block.secondary_finds:
        _add_bilingual_find_list(doc, block.secondary_finds)
    if block.likely_cause:
        if block.likely_cause.text_zh:
            doc.add_paragraph(block.likely_cause.text_zh)
        if block.likely_cause.text_en:
            doc.add_paragraph(block.likely_cause.text_en)
    if block.recommended_action:
        if block.recommended_action.text_zh:
            doc.add_paragraph(block.recommended_action.text_zh)
        if block.recommended_action.text_en:
            doc.add_paragraph(block.recommended_action.text_en)
    if block.bookmark:
        back = doc.add_paragraph()
        _add_internal_hyperlink(back, "alerts_start", "↑ Back to Alerts")


def _render_block(doc: Document, block, screenshot_fetcher: ScreenshotFetcher | None, *, include_provenance: bool = False) -> None:
    if isinstance(block, Heading):
        doc.add_heading(block.text, level=block.level)
    elif isinstance(block, Paragraph):
        doc.add_paragraph(block.text, style=block.style) if block.style else doc.add_paragraph(block.text)
    elif isinstance(block, Metadata):
        doc.add_paragraph(f"{block.label}: {block.value}")
    elif isinstance(block, Link):
        _add_link_paragraph(doc, block)
    elif isinstance(block, LogFileReference):
        _add_log_file(doc, block)
    elif isinstance(block, Screenshot):
        _add_screenshot(doc, block, screenshot_fetcher)
    elif isinstance(block, AlertNavigation):
        _add_alert_navigation(doc, block)
    elif isinstance(block, Divider):
        doc.add_paragraph("―" * 20)
    elif isinstance(block, PageBreak):
        doc.add_page_break()
    elif isinstance(block, BilingualText):
        _add_bilingual_text(doc, block)
    elif isinstance(block, BilingualFindList):
        _add_bilingual_find_list(doc, block)
    elif isinstance(block, FindList):
        _add_finds(doc, block.heading, block.finds, zh=block.language == "Chinese")
    elif isinstance(block, IncidentEvidence):
        _add_incident_evidence(doc, block, screenshot_fetcher, include_provenance=include_provenance)
    elif isinstance(block, AnalysisReference):
        _add_analysis_reference(doc, block, screenshot_fetcher, include_provenance=include_provenance)
    else:
        raise ValueError(f"render_document: unknown block type {type(block)!r}")


def render_document(
    document: ReportDocument,
    dest_path: pathlib.Path,
    screenshot_fetcher: ScreenshotFetcher | None = None,
    *,
    include_provenance: bool = False,
) -> None:
    """Generic renderer: consumes only report_document's block types,
    never a skill's own field names. Any ReportDocument — regardless of
    which skill or assembly function built it — renders through here
    unchanged."""
    doc = Document()
    _configure_document(doc, document.title, document.metadata)
    title = doc.add_heading(document.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.keep_with_next = True
    title.paragraph_format.space_after = Pt(2)
    if title.runs:
        title.runs[0].font.size = Pt(20)
        title.runs[0].font.bold = True
        title.runs[0].font.italic = True
        title.runs[0].font.color.rgb = _BODY_COLOR
    if document.metadata:
        subtitle_values = [
            item.value if item.label in {"Date", "Shift"} else f"{item.label}: {item.value}"
            for item in document.metadata
        ]
        subtitle = doc.add_paragraph(" | ".join(subtitle_values))
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(12)
        for run in subtitle.runs:
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = _MUTED_COLOR

    if include_provenance and document.provenance:
        doc.add_heading("Audit Provenance", level=2)
        for metadata in document.provenance:
            doc.add_paragraph(f"{metadata.label}: {metadata.value}")

    for block in document.blocks:
        _render_block(doc, block, screenshot_fetcher, include_provenance=include_provenance)

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest_path))


def render_daily_report_docx(
    result: dict,
    dest_path: pathlib.Path,
    screenshot_fetcher: ScreenshotFetcher | None = None,
) -> None:
    """The one skill-aware entry point: builds daily-alert-report's
    ReportDocument from its validated/merged result, then renders it
    through the generic renderer. Kept as the stable call site
    service.py uses, so this Phase 9 refactor needed no service.py
    changes."""
    render_document(build_daily_report_document(result), dest_path, screenshot_fetcher)
